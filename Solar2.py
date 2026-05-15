import os
import glob
import pandas as pd
import matplotlib.pyplot as plt
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import matplotlib
from enum import Enum
from typing import Optional, Dict, Tuple, List
from dataclasses import dataclass
from datetime import datetime

matplotlib.use('TkAgg')
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from matplotlib.figure import Figure

# ============================================================================
# ENUMS UND TYPEN
# ============================================================================

class Theme(Enum):
    """Theme-Enumeration für Type Safety."""
    LIGHT = 'light'
    DARK = 'dark'


# ============================================================================
# KONFIGURATION
# ============================================================================

@dataclass
class ThemeConfig:
    """Definiert die Farben für ein bestimmtes Theme."""
    bg_dark: str
    bg_darker: str
    bg_header: str
    fg_light: str
    color_positive: str
    color_negative: str
    matplotlib_style: str


class Config:
    """Zentrale Konfigurationsklasse für alle Konstanten."""
    
    # Pfade und Datenquellen
    DATA_DIR = '/Users/alli/Documents/01_Privat/01_Dokumente/06 Gebäude und Liegenschaft/01 Bubikon/Solaranlage Bubikon 2021/5 Energiebilanz'
    CSV_COLUMNS = [
        'Datum und Uhrzeit',
        'Gesamt Erzeugung',
        'Gesamt Verbrauch',
        'Eigenverbrauch',
        'Energie ins Netz eingespeist',
        'Energie vom Netz bezogen'
    ]
    
    # Theme-Konfigurationen (Modern & Light First)
    THEMES = {
        Theme.DARK: ThemeConfig(
            bg_dark='#1e1e1e',
            bg_darker='#2d2d2d',
            bg_header='#0f62fe',
            fg_light='#ffffff',
            color_positive='#24a148',
            color_negative='#da1e28',
            matplotlib_style='dark_background'
        ),
        Theme.LIGHT: ThemeConfig(
            bg_dark='#f4f4f4',
            bg_darker='#ffffff',
            bg_header='#0f62fe',
            fg_light='#161616',
            color_positive='#24a148',
            color_negative='#da1e28',
            matplotlib_style='default'
        )
    }
    
    # Chart-Konfiguration (Modern & Hochauflösend)
    CHART_CONFIG = {
        'figsize': (13, 6.5),
        'dpi': 120
    }
    
    CHART_STYLE = {
        'linewidth': 1.0,
        'grid_alpha': 0.25,
        'marker_size': 3
    }
    
    # UI-Konfiguration (Modern Design)
    UI_CONFIG = {
        'window_size': '1200x800',
        'window_title': 'Solar Dashboard - Solaranlage Bubikon',
        'table_font': ('Segoe UI', 11),
        'header_font': ('Segoe UI', 11),
        'title_font': ('Segoe UI', 16),
        'label_font': ('Segoe UI', 11),
        'button_font': ('Segoe UI', 11),
        'padding_large': 20,
        'padding_medium': 15,
        'padding_small': 10
    }
    
    # Monatsnamen
    MONTH_NAMES_SHORT = ['Jan', 'Feb', 'Mär', 'Apr', 'Mai', 'Jun', 'Jul', 'Aug', 'Sep', 'Okt', 'Nov', 'Dez']
    MONTH_NAMES_LONG = ['Januar', 'Februar', 'März', 'April', 'Mai', 'Juni', 
                        'Juli', 'August', 'September', 'Oktober', 'November', 'Dezember']


class ThemeManager:
    """Verwaltet das aktuelle Theme und dessen Konfiguration."""
    
    def __init__(self, initial_theme: Theme = Theme.LIGHT):
        """Initialisiere ThemeManager mit einem Theme."""
        self._current_theme = initial_theme
        self._update_matplotlib_style()
    
    @property
    def current_theme(self) -> Theme:
        """Gibt das aktuelle Theme zurück."""
        return self._current_theme
    
    @property
    def config(self) -> ThemeConfig:
        """Gibt die Konfiguration des aktuellen Themes zurück."""
        return Config.THEMES[self._current_theme]
    
    def toggle_theme(self) -> Theme:
        """Wechsle zum anderen Theme und aktualisiere Matplotlib."""
        self._current_theme = Theme.DARK if self._current_theme == Theme.LIGHT else Theme.LIGHT
        self._update_matplotlib_style()
        return self._current_theme
    
    def set_theme(self, theme: Theme) -> None:
        """Setze ein spezifisches Theme."""
        self._current_theme = theme
        self._update_matplotlib_style()
    
    def _update_matplotlib_style(self) -> None:
        """Aktualisiere Matplotlib-Style basierend auf aktuellem Theme."""
        plt.style.use(self.config.matplotlib_style)



# ============================================================================
# DATENVERARBEITUNG
# ============================================================================

class DataProcessor:
    """Verarbeitet alle Operationen mit den Solar-Daten."""
    
    @staticmethod
    def load_csv_files(data_dir: str) -> pd.DataFrame:
        """Lade und kombiniere alle CSV-Dateien aus einem Verzeichnis."""
        files = glob.glob(os.path.join(data_dir, '*.csv'))
        
        if not files:
            return pd.DataFrame()
        
        dataframes = []
        for file_path in files:
            df = pd.read_csv(file_path, skiprows=1)
            df.columns = Config.CSV_COLUMNS
            
            # Datum konvertieren und Zeit-Felder extrahieren
            df['Datum und Uhrzeit'] = pd.to_datetime(df['Datum und Uhrzeit'], format='%d.%m.%Y')
            df['Jahr'] = df['Datum und Uhrzeit'].dt.year
            df['Woche'] = df['Datum und Uhrzeit'].dt.isocalendar()['week']
            df['Monat'] = df['Datum und Uhrzeit'].dt.month
            
            dataframes.append(df)
        
        return pd.concat(dataframes, ignore_index=True) if dataframes else pd.DataFrame()
    
    @staticmethod
    def convert_to_kwh(pivot_df: pd.DataFrame) -> pd.DataFrame:
        """Konvertiere Wh zu kWh (Division durch 1000)."""
        return pivot_df / 1000
    
    @staticmethod
    def create_pivot_table(summary_df: pd.DataFrame, index_name: str) -> pd.DataFrame:
        """Erstelle Pivot-Tabelle (generisch für Wochen und Monate)."""
        pivot_df = summary_df.pivot(index=index_name, columns='Jahr', values='Gesamt Erzeugung')
        return DataProcessor.convert_to_kwh(pivot_df)
    
    @staticmethod
    def calculate_weekly_sums(df: pd.DataFrame) -> pd.DataFrame:
        """Berechne Summen der Gesamt Erzeugung pro Woche und Jahr."""
        return df.groupby(['Woche', 'Jahr'])['Gesamt Erzeugung'].sum().reset_index()
    
    @staticmethod
    def calculate_monthly_sums(df: pd.DataFrame) -> pd.DataFrame:
        """Berechne Summen der Gesamt Erzeugung pro Monat und Jahr."""
        return df.groupby(['Monat', 'Jahr'])['Gesamt Erzeugung'].sum().reset_index()
    
    @staticmethod
    def calculate_running_total(pivot_df: pd.DataFrame) -> pd.DataFrame:
        """Berechne kumulatives Running Total pro Jahr."""
        return pivot_df.cumsum()
    
    @staticmethod
    def extract_year_comparison(df: pd.DataFrame) -> Optional[Tuple[int, int]]:
        """Extrahiere die letzten zwei Jahre aus einem DataFrame für Vergleiche."""
        if len(df.columns) < 2:
            return None
        years = sorted(df.columns)[-2:]
        return years[0], years[1]
    
    @staticmethod
    def calculate_deviations(running_total_df: pd.DataFrame, index_name: str = 'Woche') -> Optional[Tuple]:
        """Berechne Abweichungen zwischen aktuellem und Vorjahr."""
        result = DataProcessor.extract_year_comparison(running_total_df)
        if result is None:
            return None
        
        prev_year, current_year = result
        deviation = running_total_df[current_year] - running_total_df[prev_year]
        
        return pd.DataFrame({
            index_name: running_total_df.index,
            'Abweichung': deviation.values
        }), current_year, prev_year
    
    @staticmethod
    def create_daily_data_with_metrics(df_raw: pd.DataFrame) -> pd.DataFrame:
        """Erstelle aggregierte Tagesdaten mit allen Kennzahlen."""
        df_daily = df_raw.copy()
        df_daily['Datum'] = df_daily['Datum und Uhrzeit'].dt.date
        
        # Aggregiere auf Tagesbasis
        daily_agg = df_daily.groupby('Datum').agg({
            'Gesamt Erzeugung': 'sum',
            'Gesamt Verbrauch': 'sum',
            'Eigenverbrauch': 'sum',
            'Energie ins Netz eingespeist': 'sum',
            'Energie vom Netz bezogen': 'sum'
        }).reset_index()
        
        # Konvertiere Wh zu kWh
        for col in daily_agg.columns[1:]:
            daily_agg[col] = (daily_agg[col] / 1000).round(2)
        
        # Berechne zusätzliche Kennzahlen
        daily_agg['Eigenverbrauchsquote [%]'] = (
            (daily_agg['Eigenverbrauch'] / daily_agg['Gesamt Erzeugung'] * 100)
            .fillna(0).round(2)
        )
        
        daily_agg['Autarkie [%]'] = (
            (daily_agg['Eigenverbrauch'] / (daily_agg['Eigenverbrauch'] + daily_agg['Energie vom Netz bezogen']) * 100)
            .fillna(0).round(2)
        )
        
        # Spalten umbenennen
        daily_agg.columns = [
            'Datum',
            'Gesamt Erzeugung [kWh]',
            'Gesamt Verbrauch [kWh]',
            'Eigenverbrauch [kWh]',
            'Energie ins Netz [kWh]',
            'Energie aus Netz [kWh]',
            'Eigenverbrauchsquote [%]',
            'Autarkie [%]'
        ]
        
        return daily_agg.sort_values('Datum', ascending=False)

# ============================================================================
# GRAFIKRENDERING
# ============================================================================

class ChartRenderer:
    """Rendert alle Arten von Grafiken mit konsistentem Styling."""
    
    def __init__(self, theme_manager: ThemeManager):
        """Initialisiere ChartRenderer mit Theme-Manager."""
        self.theme_manager = theme_manager
    
    def create_line_chart(self, frame: tk.Frame, data: pd.DataFrame, title: str,
                         xlabel: str, ylabel: str) -> None:
        """Erstelle ein modernes Liniendiagramm mit besserer Ästhetik."""
        fig = Figure(**Config.CHART_CONFIG)
        fig.patch.set_facecolor(self.theme_manager.config.bg_darker)
        ax = fig.add_subplot(111)
        
        # Moderne Farben für verschiedene Jahre
        colors = ['#0f62fe', '#24a148', '#f1c21b', '#ff832b', '#da1e28']
        
        for idx, year in enumerate(data.columns):
            color = colors[idx % len(colors)]
            ax.plot(data.index, data[year], linestyle='solid', linewidth=Config.CHART_STYLE['linewidth'],
                   label=str(year), marker='o', markersize=Config.CHART_STYLE['marker_size'],
                   color=color, alpha=0.85)
        
        # X-Achsen-Ticks dynamisch setzen
        all_ticks = sorted(data.index.tolist())
        if all_ticks:
            x_ticks = all_ticks[::2] if len(all_ticks) > 26 else all_ticks
            ax.set_xticks(x_ticks)
            ax.set_xlim(all_ticks[0] - 0.5, all_ticks[-1] + 0.5)
        
        ax.set_xlabel(xlabel, fontsize=11)
        ax.set_ylabel(ylabel, fontsize=11)
        ax.set_title(title, fontsize=13, pad=15)
        ax.legend(title='PV-Jahr', fontsize=10, title_fontsize=11, loc='best', framealpha=0.95)
        ax.grid(linestyle='--', alpha=Config.CHART_STYLE['grid_alpha'], linewidth=0.6)
        ax.tick_params(axis='x', rotation=45, labelsize=9)
        ax.tick_params(axis='y', labelsize=9)
        
        # Moderne Achsen-Styling
        ax.spines['top'].set_visible(False)
        ax.spines['right'].set_visible(False)
        
        fig.tight_layout()
        
        canvas = FigureCanvasTkAgg(fig, master=frame)
        canvas.draw()
        canvas.get_tk_widget().pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
    
    def create_weekly_deviation_chart(self, frame: tk.Frame, deviations_df: pd.DataFrame) -> None:
        """Erstelle moderne Balkengrafik der wöchentlichen Abweichungen."""
        data = deviations_df.sort_values('Woche').reset_index(drop=True)
        weeks = data['Woche'].values
        
        # Intelligente Tick-Auswahl
        if len(weeks) > 26:
            x_ticks = weeks[::2] if len(weeks) % 2 == 0 else list(weeks[::2]) + [weeks[-1]]
        else:
            x_ticks = weeks
        
        # Moderne Farben basierend auf positiv/negativ
        colors = [
            self.theme_manager.config.color_positive if dev >= 0 else self.theme_manager.config.color_negative
            for dev in data['Abweichung'].values
        ]
        
        fig = Figure(**Config.CHART_CONFIG)
        fig.patch.set_facecolor(self.theme_manager.config.bg_darker)
        ax = fig.add_subplot(111)
        
        bars = ax.bar(weeks, data['Abweichung'], color=colors, edgecolor='none',
                     width=0.7, alpha=0.85)
        
        ax.set_xticks(x_ticks)
        if len(weeks) > 0:
            ax.set_xlim(weeks.min() - 0.5, weeks.max() + 0.5)
        
        ax.set_xlabel('Kalenderwoche', fontsize=11)
        ax.set_ylabel('Abweichung [kWh]', fontsize=11)
        ax.set_title('Abweichung des kumulativen Totals: Aktuelles Jahr vs. Vorjahr',
                    fontsize=13, pad=15)
        ax.axhline(y=0, color='#525252', linestyle='-', linewidth=1.2, alpha=0.6)
        ax.grid(linestyle='--', alpha=Config.CHART_STYLE['grid_alpha'], axis='y', linewidth=0.6)
        ax.tick_params(axis='x', rotation=45, labelsize=9)
        ax.tick_params(axis='y', labelsize=9)
        
        # Moderne Achsen-Styling
        ax.spines['top'].set_visible(False)
        ax.spines['right'].set_visible(False)
        ax.spines['left'].set_linewidth(0.7)
        ax.spines['bottom'].set_linewidth(0.7)
        
        fig.tight_layout()
        
        canvas = FigureCanvasTkAgg(fig, master=frame)
        canvas.draw()
        canvas.get_tk_widget().pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
    
    def create_monthly_deviation_chart(self, frame: tk.Frame, deviations_df: pd.DataFrame) -> None:
        """Erstelle moderne Balkengrafik der monatlichen Abweichungen mit Wertlabeln."""
        data = deviations_df.sort_values('Monat').reset_index(drop=True)
        months = data['Monat'].values
        month_labels = [Config.MONTH_NAMES_SHORT[int(m)-1] for m in months]
        
        # Moderne Farben basierend auf positiv/negativ
        colors = [
            self.theme_manager.config.color_positive if dev >= 0 else self.theme_manager.config.color_negative
            for dev in data['Abweichung'].values
        ]
        
        fig = Figure(**Config.CHART_CONFIG)
        fig.patch.set_facecolor(self.theme_manager.config.bg_darker)
        ax = fig.add_subplot(111)
        
        bars = ax.bar(range(len(months)), data['Abweichung'], color=colors,
                     edgecolor='none', width=0.65, alpha=0.85)
        
        # Wertlabels direkt auf Balken hinzufügen (nah an den Säulen)
        label_color = '#161616' if self.theme_manager.current_theme == Theme.LIGHT else '#ffffff'
        for bar, value in zip(bars, data['Abweichung'].values):
            height = bar.get_height()
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                height + (2 if height >= 0 else -2),
                f'{value:.0f}',
                ha='center',
                va='bottom' if height >= 0 else 'top',
                fontsize=10,
                color=label_color
            )
        
        ax.set_xticks(range(len(months)))
        ax.set_xticklabels(month_labels, fontsize=10)
        ax.set_xlabel('Monat', fontsize=11)
        ax.set_ylabel('Abweichung [kWh]', fontsize=11)
        ax.set_title('Abweichung monatlicher Summen: Aktuelles Jahr vs. Vorjahr',
                    fontsize=13, pad=15)
        ax.axhline(y=0, color='#525252', linestyle='-', linewidth=1.2, alpha=0.6)
        ax.grid(linestyle='--', alpha=Config.CHART_STYLE['grid_alpha'], axis='y', linewidth=0.6)
        ax.tick_params(axis='y', labelsize=9)
        
        # Moderne Achsen-Styling
        ax.spines['top'].set_visible(False)
        ax.spines['right'].set_visible(False)
        ax.spines['left'].set_linewidth(0.7)
        ax.spines['bottom'].set_linewidth(0.7)
        
        fig.tight_layout()
        
        canvas = FigureCanvasTkAgg(fig, master=frame)
        canvas.draw()
        canvas.get_tk_widget().pack(fill=tk.BOTH, expand=True, padx=10, pady=10)


# ============================================================================
# TABELLENRENDERING
# ============================================================================

class TableRenderer:
    """Rendert moderne, gut lesbare Tabellen mit konsistentem Styling."""
    
    def __init__(self, theme_manager: ThemeManager):
        """Initialisiere TableRenderer mit Theme-Manager."""
        self.theme_manager = theme_manager
    
    def create_table(self, parent_frame: tk.Frame, table_data: pd.DataFrame, title: str) -> None:
        """Erstelle moderne, scrollbare Tabelle mit optimalem Styling."""
        # Titel mit besserer Formatierung
        title_label = ttk.Label(parent_frame, text=title, font=Config.UI_CONFIG['title_font'])
        title_label.pack(pady=(15, 20), padx=15)
        
        # Header mit modernem Design
        columns = list(table_data.columns)
        header_frame = tk.Frame(parent_frame, bg=self.theme_manager.config.bg_header, height=40)
        header_frame.pack(fill=tk.X, padx=15, pady=(0, 0))
        header_frame.pack_propagate(False)
        
        for col_name in columns:
            label = tk.Label(
                header_frame,
                text=col_name,
                font=Config.UI_CONFIG['header_font'],
                bg=self.theme_manager.config.bg_header,
                fg='#ffffff',
                anchor='center',
                padx=10,
                pady=8
            )
            label.pack(side=tk.LEFT, expand=True, fill=tk.BOTH)
        
        # Treeview-Container mit Scrollbar
        container_frame = ttk.Frame(parent_frame)
        container_frame.pack(fill=tk.BOTH, expand=True, padx=15, pady=(5, 15))
        
        scrollbar = ttk.Scrollbar(container_frame)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        tree = ttk.Treeview(
            container_frame,
            columns=columns,
            show='tree',
            yscrollcommand=scrollbar.set,
            height=22
        )
        scrollbar.config(command=tree.yview)
        
        # Spalten konfigurieren (optimal dimensioniert)
        tree.column('#0', width=0, stretch=tk.NO)
        for col in columns:
            tree.column(col, anchor='center', width=130)
        
        # Daten mit modernem Zeilen-Styling einfügen
        for idx, row_data in enumerate(table_data.values):
            tag = 'oddrow' if idx % 2 else 'evenrow'
            tree.insert('', 'end', values=[f'{val:.2f}' if isinstance(val, float) else str(val) 
                                           for val in row_data], tags=(tag,))
        
        # Modernes Styling für Zeilen (Light Theme optimiert)
        if self.theme_manager.current_theme == Theme.LIGHT:
            tree.tag_configure('oddrow', background='#f4f4f4', foreground='#161616')
            tree.tag_configure('evenrow', background='#ffffff', foreground='#161616')
        else:
            tree.tag_configure('oddrow', background='#2d2d2d', foreground='#ffffff')
            tree.tag_configure('evenrow', background='#1e1e1e', foreground='#ffffff')
        
        tree.pack(fill=tk.BOTH, expand=True)


# ============================================================================
# EXPORT-FUNKTIONEN
# ============================================================================

def export_to_excel(df_raw: pd.DataFrame, parent_window: Optional[tk.Widget] = None) -> None:
    """Exportiere tägliche Daten mit Kennzahlen zu Excel-Datei."""
    try:
        # Datei-Dialog öffnen
        file_path = filedialog.asksaveasfilename(
            title='Excel-Datei speichern',
            defaultextension='.xlsx',
            filetypes=[('Excel Dateien', '*.xlsx'), ('Alle Dateien', '*.*')],
            initialfile=f"Solaranlage_Tagesbilanz_{datetime.now().strftime('%Y%m%d')}.xlsx"
        )
        
        if not file_path:
            return  # Benutzer hat Abbruch geklickt
        
        # Tägliche Daten mit Kennzahlen erstellen
        daily_df = DataProcessor.create_daily_data_with_metrics(df_raw)
        
        # Zu Excel exportieren mit Formatierung
        with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
            daily_df.to_excel(writer, sheet_name='Tagesbilanz', index=False)
            
            # Formatierung anwenden
            workbook = writer.book
            worksheet = writer.sheets['Tagesbilanz']
            
            # Spaltenbreiten anpassen
            worksheet.column_dimensions['A'].width = 12
            for col in range(2, len(daily_df.columns) + 1):
                worksheet.column_dimensions[chr(64 + col)].width = 16
            
            # Header formatieren
            from openpyxl.styles import Font, PatternFill, Alignment
            header_fill = PatternFill(start_color='1976D2', end_color='1976D2', fill_type='solid')
            header_font = Font(bold=True, color='FFFFFF', size=11)
            
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
            
            # Datenzellenformatierung (Zentrieren)
            for row in worksheet.iter_rows(min_row=2, max_row=len(daily_df) + 1, 
                                           min_col=1, max_col=len(daily_df.columns)):
                for cell in row:
                    cell.alignment = Alignment(horizontal='center', vertical='center')
        
        messagebox.showinfo('Erfolg', f'Datei erfolgreich gespeichert:\n{file_path}')
        
    except ImportError:
        messagebox.showerror('Fehler', 'openpyxl-Paket nicht installiert. '
                            'Bitte installieren Sie: pip install openpyxl')
    except Exception as e:
        messagebox.showerror('Fehler', f'Fehler beim Exportieren:\n{str(e)}')


# ============================================================================
# HAUPTANWENDUNG
# ============================================================================

class SolarDashboard:
    """Orchestriert die gesamte GUI-Anwendung."""
    
    def __init__(self, pivot_weekly: pd.DataFrame, pivot_monthly: pd.DataFrame, df_raw: pd.DataFrame):
        """Initialisiere das Solar Dashboard mit Daten."""
        self.pivot_weekly = pivot_weekly
        self.pivot_monthly = pivot_monthly
        self.df_raw = df_raw
        
        # Manager und Renderer
        self.theme_manager = ThemeManager(Theme.LIGHT)
        self.chart_renderer = ChartRenderer(self.theme_manager)
        self.table_renderer = TableRenderer(self.theme_manager)
        
        # Hauptfenster
        self.root = tk.Tk()
        self.root.title(Config.UI_CONFIG['window_title'])
        self.root.geometry(Config.UI_CONFIG['window_size'])
        
        # GUI-State
        self.toolbar = None
        self.notebook = None
        self.theme_var = None
    
    def _setup_toolbar(self) -> None:
        """Erstelle die moderne Werkzeugleiste mit Theme-Schalter."""
        self.toolbar = tk.Frame(self.root, bg=self.theme_manager.config.bg_header, height=50)
        self.toolbar.pack(fill=tk.X, padx=0, pady=0, side=tk.TOP)
        self.toolbar.pack_propagate(False)
        
        # Logo/Titel Section
        title_frame = tk.Frame(self.toolbar, bg=self.theme_manager.config.bg_header)
        title_frame.pack(side=tk.LEFT, padx=20, pady=10, fill=tk.BOTH, expand=False)
        
        title_label = tk.Label(
            title_frame,
            text='☀️ Solar Dashboard',
            font=('Segoe UI', 13),
            bg=self.theme_manager.config.bg_header,
            fg='#ffffff'
        )
        title_label.pack(side=tk.LEFT)
        
        # Spacer
        spacer = tk.Frame(self.toolbar, bg=self.theme_manager.config.bg_header)
        spacer.pack(side=tk.LEFT, expand=True, fill=tk.X)
        
        # Theme Schalter Section (rechts)
        control_frame = tk.Frame(self.toolbar, bg=self.theme_manager.config.bg_header)
        control_frame.pack(side=tk.RIGHT, padx=20, pady=8, fill=tk.BOTH)
        
        theme_label = tk.Label(
            control_frame,
            text='Design:',
            font=Config.UI_CONFIG['label_font'],
            bg=self.theme_manager.config.bg_header,
            fg='#ffffff'
        )
        theme_label.pack(side=tk.LEFT, padx=(0, 10))
        
        # Checkbox für Helles Theme (modern gestyled)
        self.theme_var = tk.BooleanVar(value=(self.theme_manager.current_theme == Theme.LIGHT))
        checkbox = tk.Checkbutton(
            control_frame,
            text='Helles Design',
            variable=self.theme_var,
            command=self._toggle_theme,
            bg=self.theme_manager.config.bg_header,
            fg='#ffffff',
            font=Config.UI_CONFIG['label_font'],
            selectcolor=self.theme_manager.config.bg_header,
            activebackground=self.theme_manager.config.bg_header,
            activeforeground='#ffffff',
            highlightthickness=0,
            bd=0,
            padx=5
        )
        checkbox.pack(side=tk.LEFT)
    
    
    def _toggle_theme(self) -> None:
        """Wechsle das Theme und aktualisiere nur die Farben (nicht die ganze GUI)."""
        is_light = self.theme_var.get()
        new_theme = Theme.LIGHT if is_light else Theme.DARK
        self.theme_manager.set_theme(new_theme)
        
        # Aktualisiere nur Farben, nicht komplette GUI
        self._update_theme_colors()
    
    def _update_theme_colors(self) -> None:
        """Aktualisiere Fenster-Farben für das neue Theme."""
        # Hauptfenster
        self.root.configure(bg=self.theme_manager.config.bg_dark)
        
        # Toolbar
        if self.toolbar:
            self.toolbar.configure(bg=self.theme_manager.config.bg_header)
            for widget in self.toolbar.winfo_children():
                if isinstance(widget, tk.Label):
                    widget.configure(
                        bg=self.theme_manager.config.bg_header,
                        fg=self.theme_manager.config.fg_light
                    )
                elif isinstance(widget, tk.Checkbutton):
                    widget.configure(
                        bg=self.theme_manager.config.bg_header,
                        fg=self.theme_manager.config.fg_light,
                        selectcolor=self.theme_manager.config.bg_header,
                        activebackground=self.theme_manager.config.bg_header,
                        activeforeground=self.theme_manager.config.fg_light
                    )
    
    def _setup_theme_style(self) -> None:
        """Konfiguriere das ttk-Theme."""
        style = ttk.Style()
        style.theme_use('clam')
        
        # Notebook Style
        style.configure('TNotebook', background=self.theme_manager.config.bg_dark, borderwidth=0)
        style.configure('TNotebook.Tab', padding=[10, 10],
                       background=self.theme_manager.config.bg_darker,
                       foreground=self.theme_manager.config.fg_light)
        style.map('TNotebook.Tab',
                 background=[('selected', self.theme_manager.config.bg_header)],
                 foreground=[('selected', '#ffffff')])
        
        # Frame, Label und Treeview Styles
        style.configure('TFrame', background=self.theme_manager.config.bg_dark,
                       foreground=self.theme_manager.config.fg_light)
        style.configure('TLabel', background=self.theme_manager.config.bg_dark,
                       foreground=self.theme_manager.config.fg_light)
        style.configure('Treeview', background=self.theme_manager.config.bg_darker,
                       foreground=self.theme_manager.config.fg_light,
                       fieldbackground=self.theme_manager.config.bg_darker)
        style.map('Treeview', background=[('selected', self.theme_manager.config.bg_header)])
    
    def _create_tabs(self) -> None:
        """Erstelle alle Reiter der Anwendung."""
        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        # Reiter erstellen
        self._create_tab_overview()
        self._create_tab_weekly_data()
        self._create_tab_running_total()
        self._create_tab_weekly_deviations()
        self._create_tab_monthly_data()
        self._create_tab_monthly_deviations()
        self._create_tab_control()
        
        # Standardmäßig beim Übersicht-Reiter starten
        self.notebook.select(0)
    
    def _create_tab_overview(self) -> None:
        """Reiter 1: Liniendiagramm wöchentlicher Erzeugung."""
        frame = ttk.Frame(self.notebook)
        self.notebook.add(frame, text='Übersicht')
        self.chart_renderer.create_line_chart(
            frame, self.pivot_weekly,
            'Wöchentliche Gesamt-Erzeugung im Jahresvergleich',
            'Kalenderwoche', 'Gesamt Erzeugung [kWh]'
        )
    
    def _create_tab_weekly_data(self) -> None:
        """Reiter 2: Tabelle mit wöchentlichen Daten."""
        frame = ttk.Frame(self.notebook)
        self.notebook.add(frame, text='Wochendaten')
        table_data = self.pivot_weekly.round(1).reset_index()
        self.table_renderer.create_table(frame, table_data,
                                        'Wöchentliche Gesamt-Erzeugung [kWh]')
    
    def _create_tab_running_total(self) -> None:
        """Reiter 3: Kumulatives Running Total."""
        frame = ttk.Frame(self.notebook)
        self.notebook.add(frame, text='Running Total')
        running_total = DataProcessor.calculate_running_total(self.pivot_weekly)
        table_data = running_total.round(1).reset_index()
        self.table_renderer.create_table(frame, table_data,
                                        'Kumulatives Running Total pro Jahr [kWh]')
    
    def _create_tab_weekly_deviations(self) -> None:
        """Reiter 4: Wöchentliche Abweichungen."""
        frame = ttk.Frame(self.notebook)
        self.notebook.add(frame, text='Abweichungen')
        
        running_total = DataProcessor.calculate_running_total(self.pivot_weekly)
        result = DataProcessor.calculate_deviations(running_total, 'Woche')
        
        if result is not None:
            deviations_df, _, _ = result
            self.chart_renderer.create_weekly_deviation_chart(frame, deviations_df)
    
    def _create_tab_monthly_data(self) -> None:
        """Reiter 5: Tabelle mit monatlichen Daten."""
        frame = ttk.Frame(self.notebook)
        self.notebook.add(frame, text='Monatsdaten')
        
        table_data = self.pivot_monthly.round(1).reset_index()
        
        # Monatsnamen hinzufügen
        table_data.insert(1, 'Monatname',
                         table_data['Monat'].apply(lambda x: Config.MONTH_NAMES_LONG[int(x)-1]))
        table_data = table_data.drop('Monat', axis=1)
        table_data = table_data.rename(columns={'Monatname': 'Monat'})
        
        self.table_renderer.create_table(frame, table_data,
                                        'Monatliche Gesamt-Erzeugung [kWh]')
    
    def _create_tab_monthly_deviations(self) -> None:
        """Reiter 6: Monatliche Abweichungen."""
        frame = ttk.Frame(self.notebook)
        self.notebook.add(frame, text='Monatliche Abweichungen')
        
        result = DataProcessor.calculate_deviations(self.pivot_monthly, 'Monat')
        
        if result is not None:
            deviations_df, _, _ = result
            self.chart_renderer.create_monthly_deviation_chart(frame, deviations_df)
    
    def _create_tab_control(self) -> None:
        """Reiter 7: Moderne Steuerung und Export-Funktionen."""
        frame = ttk.Frame(self.notebook)
        self.notebook.add(frame, text='Steuerung')
        
        # Hauptcontainer mit modernem Design
        main_container = ttk.Frame(frame, padding=Config.UI_CONFIG['padding_large'])
        main_container.pack(fill=tk.BOTH, expand=True)
        
        # Titel
        title_label = ttk.Label(main_container, text='📊 Datenexport & Steuerung',
                               font=Config.UI_CONFIG['title_font'])
        title_label.pack(pady=(0, Config.UI_CONFIG['padding_large']), anchor='w')
        
        # Export-Sektion (moderne Gestaltung)
        export_frame = ttk.LabelFrame(main_container, text='📁 Excel Export', 
                                     padding=Config.UI_CONFIG['padding_medium'])
        export_frame.pack(fill=tk.X, pady=Config.UI_CONFIG['padding_medium'])
        
        export_button = tk.Button(
            export_frame,
            text='💾 Tagesbilanz exportieren',
            command=lambda: export_to_excel(self.df_raw, frame),
            bg='#2a2a2a',
            fg='#000000',
            font=Config.UI_CONFIG['button_font'],
            padx=20,
            pady=12,
            relief=tk.FLAT,
            activebackground='#0f62fe',
            activeforeground='#ffffff',
            cursor='hand2',
            bd=0
        )
        export_button.pack(side=tk.LEFT, padx=Config.UI_CONFIG['padding_medium'])
        
        info_text = ttk.Label(
            export_frame,
            text='Exportiert die Tagesbilanz mit allen Kennzahlen als Excel-Datei (Format: .xlsx)',
            font=Config.UI_CONFIG['label_font'],
            wraplength=600
        )
        info_text.pack(side=tk.LEFT, padx=Config.UI_CONFIG['padding_medium'], fill=tk.BOTH, expand=True)
        
        # Info-Sektion (modern & übersichtlich)
        info_frame = ttk.LabelFrame(main_container, text='ℹ️ Über diese Anwendung',
                                   padding=Config.UI_CONFIG['padding_medium'])
        info_frame.pack(fill=tk.BOTH, expand=True, pady=Config.UI_CONFIG['padding_medium'])
        
        info_label = ttk.Label(
            info_frame,
            text=f'''📈 Solar Dashboard - Analyse & Visualisierung der Solaranlage Bubikon

📁 Datenquelle:
   {Config.DATA_DIR}

📊 Verfügbare Reiter & Funktionen:
   🔹 Übersicht        → Wöchentliche Erzeugung im Jahresvergleich (Liniendiagramm)
   🔹 Wochendaten      → Detaillierte wöchentliche Daten-Tabelle
   🔹 Running Total    → Kumulierte Erzeugung pro Woche
   🔹 Abweichungen     → Wöchentliche Abweichungen (aktuelles Jahr vs. Vorjahr)
   🔹 Monatsdaten      → Monatliche Erzeugung im Jahresvergleich
   🔹 Monatliche Abw.  → Monatliche Abweichungen mit Wertlabeln
   🔹 Steuerung        → Datenexport und Anwendungsinfo

💡 Hinweise & Erklärungen:
   • Alle Energiewerte sind in kWh (Kilowattstunden) angegeben
   • Eigenverbrauchsquote: Anteil der erzeugten Energie, der vor Ort verbraucht wird
   • Autarkie: Grad der Unabhängigkeit vom öffentlichen Stromnetz
   • Excel-Export enthält Tagesbilanz mit zusätzlichen Kennzahlen
   • Design: Helles/Dunkles Theme über den Regler in der Kopfzeile wählbar

✨ Moderne Visualisierung mit optimierter Lesbarkeit und Farbcodierung
''',
            font=Config.UI_CONFIG['label_font'],
            justify=tk.LEFT
        )
        info_label.pack(fill=tk.BOTH, expand=True)

    
    def run(self) -> None:
        """Starten Sie die Anwendung."""
        # Fenster konfigurieren
        self.root.configure(bg=self.theme_manager.config.bg_dark)
        
        # GUI aufbauen
        self._setup_toolbar()
        self._setup_theme_style()
        self._create_tabs()
        
        # Fenster starten
        self.root.mainloop()


if __name__ == '__main__':
    # Daten laden
    df = DataProcessor.load_csv_files(Config.DATA_DIR)
    
    if df.empty:
        print(f"Fehler: Keine CSV-Dateien im Verzeichnis gefunden: {Config.DATA_DIR}")
        exit(1)
    
    # Pivot-Tabellen erstellen
    weekly_sums = DataProcessor.calculate_weekly_sums(df)
    pivot_weekly = DataProcessor.create_pivot_table(weekly_sums, 'Woche')
    
    monthly_sums = DataProcessor.calculate_monthly_sums(df)
    pivot_monthly = DataProcessor.create_pivot_table(monthly_sums, 'Monat')
    
    # Dashboard starten
    dashboard = SolarDashboard(pivot_weekly, pivot_monthly, df)
    dashboard.run()
